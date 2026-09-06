"""`adapters/documents.py` + `POST /api/v1/documents/extract`.

Input 노드가 PDF·DOCX·텍스트 파일에서 본문을 받아 오는 경로다. 픽스처는
**실제 파일 바이트**를 그 자리에서 만든다 — 리포에 바이너리를 넣지 않으면서
라이브러리를 진짜로 태우기 위함이다(모킹하면 pdfplumber/python-docx 업그레이드가
깨져도 이 테스트가 못 잡는다).
"""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient

from app.adapters import documents as docs
from app.main import app

client = TestClient(app)


# ────────────────────────────── 픽스처 ──────────────────────────────


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                t.rows[r].cells[c].text = cell
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf(pages: list[str]) -> bytes:
    import fitz

    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((60, 80), text)
    return doc.tobytes()


# ────────────────────────────── 어댑터 ──────────────────────────────


def test_pdf_text_and_page_count():
    result = docs.extract_document("a.pdf", make_pdf(["First page.", "Second page."]))
    assert result.pages == 2
    assert "First page." in result.text
    assert "Second page." in result.text
    assert result.truncated is False


def test_docx_paragraphs_and_tables():
    data = make_docx(["첫 문단", "둘째 문단"], table=[["키", "값"]])
    result = docs.extract_document("a.docx", data)
    assert "첫 문단" in result.text
    assert "둘째 문단" in result.text
    # 표 안의 글자도 원문의 일부다 — 빼면 사용자는 "내용이 사라졌다"로 본다.
    assert "키\t값" in result.text
    assert result.pages is None


def test_line_breaks_are_preserved():
    """Input 노드가 줄바꿈을 그대로 받아야 한다 — 이게 이 기능의 요점이다."""
    result = docs.extract_document("a.txt", "one\ntwo\n\nthree".encode())
    assert result.text == "one\ntwo\n\nthree"


def test_excess_blank_lines_collapse_but_paragraphs_survive():
    result = docs.extract_document("a.txt", "a\n\n\n\n\nb".encode())
    # 문단 구분(빈 줄 1개)은 원문의 정보라 남고, 그 이상만 접힌다.
    assert result.text == "a\n\nb"


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16", "cp949"])
def test_text_decoding_fallbacks(encoding):
    result = docs.extract_document("a.txt", "한글 본문".encode(encoding))
    assert "한글 본문" in result.text


def test_unsupported_extension():
    with pytest.raises(docs.UnsupportedDocumentError):
        docs.extract_document("a.exe", b"MZ")


def test_extensionless_file_is_unsupported():
    with pytest.raises(docs.UnsupportedDocumentError):
        docs.extract_document("README", b"hello")


def test_too_large():
    with pytest.raises(docs.DocumentTooLargeError):
        docs.extract_document("a.txt", b"x" * (docs.MAX_UPLOAD_BYTES + 1))


def test_corrupt_pdf_raises_parse_error():
    with pytest.raises(docs.DocumentParseError):
        docs.extract_document("a.pdf", b"definitely not a pdf")


def test_truncation_is_reported():
    data = ("y" * (docs.MAX_TEXT_CHARS + 500)).encode()
    result = docs.extract_document("a.txt", data)
    assert result.truncated is True
    assert result.chars == docs.MAX_TEXT_CHARS


# ────────────────────────────── 라우터 ──────────────────────────────


def test_route_extracts_pdf():
    res = client.post(
        "/api/v1/documents/extract",
        files={"file": ("sample.pdf", make_pdf(["Hello from a PDF."]), "application/pdf")},
    )
    assert res.status_code == 200
    body = res.json()
    assert body["pages"] == 1
    assert "Hello from a PDF." in body["text"]
    assert body["filename"] == "sample.pdf"


def test_route_extracts_docx():
    res = client.post(
        "/api/v1/documents/extract",
        files={"file": ("sample.docx", make_docx(["본문입니다"]), "application/octet-stream")},
    )
    assert res.status_code == 200
    assert "본문입니다" in res.json()["text"]


def test_route_unsupported_returns_ac_e406():
    """트리거: AC-E406"""
    res = client.post("/api/v1/documents/extract", files={"file": ("x.exe", b"MZ")})
    assert res.status_code == 415
    assert res.json()["error"]["code"] == "AC-E406"


def test_route_too_large_returns_ac_e407():
    """트리거: AC-E407"""
    res = client.post(
        "/api/v1/documents/extract",
        files={"file": ("big.txt", b"x" * (docs.MAX_UPLOAD_BYTES + 1))},
    )
    assert res.status_code == 413
    assert res.json()["error"]["code"] == "AC-E407"


def test_route_corrupt_returns_ac_e408():
    """트리거: AC-E408"""
    res = client.post("/api/v1/documents/extract", files={"file": ("bad.pdf", b"nope")})
    assert res.status_code == 422
    assert res.json()["error"]["code"] == "AC-E408"


def test_route_rejects_pdf_without_text_layer():
    """스캔 PDF 는 '성공 + 빈 문자열'로 온다. 조용히 빈 Input 을 만들면 안 된다."""
    import fitz

    doc = fitz.open()
    doc.new_page()  # 글자 없는 빈 페이지
    res = client.post("/api/v1/documents/extract", files={"file": ("scan.pdf", doc.tobytes())})
    assert res.status_code == 422
    assert res.json()["error"]["code"] == "AC-E408"


def test_route_never_persists_the_upload():
    """파일은 서버 어디에도 남지 않는다 (라우터 docstring 의 계약)."""
    from pathlib import Path

    from app.config import get_settings

    workspace = Path(get_settings().workspace_dir)
    before = set(p.name for p in workspace.rglob("*")) if workspace.exists() else set()
    client.post(
        "/api/v1/documents/extract",
        files={"file": ("secret.txt", b"confidential contents")},
    )
    after = set(p.name for p in workspace.rglob("*")) if workspace.exists() else set()
    assert before == after
