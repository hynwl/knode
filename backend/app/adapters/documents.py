"""업로드한 문서(PDF · DOCX · 평문)에서 텍스트를 뽑는다.

Input 노드에 긴 원문을 손으로 붙여넣는 대신 파일을 집어넣을 수 있게 하기 위한
것이다. 추출한 **텍스트만** 노드로 돌아가고, 파일 자체는 서버 어디에도 남지
않는다 — 요청 처리 동안 메모리에만 있다가 버려진다.

## 왜 백엔드에서 하나

브라우저에서 PDF/DOCX 를 파싱하려면 pdf.js·mammoth 같은 무거운 번들을 프론트에
추가해야 하는데, 백엔드에는 `crewai-tools` 가 끌고 온 `pdfplumber` 와
`python-docx` 가 **이미** 있다. 다만 여기서 직접 import 하므로
`requirements.txt` 규칙대로 명시 고정한다.

## 한도

`MAX_UPLOAD_BYTES` 와 `MAX_TEXT_CHARS` 두 겹으로 막는다. 인증이 없는 백엔드라
(§19.2) 큰 파일 하나로 메모리를 밀어내는 걸 막아야 하고, 추출된 텍스트는 결국
LLM 프롬프트에 들어가므로 사용자가 의도치 않게 수십만 자를 태우는 것도 막는다.
자르면 자른 사실을 `truncated` 로 알려 준다 — 조용히 삼키면 사용자는 앞부분만
반영된 결과를 보고도 이유를 모른다.
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass

logger = logging.getLogger(__name__)

#: 업로드 상한. 문서 하나를 프롬프트에 넣는 용도라 넉넉할 이유가 없다.
MAX_UPLOAD_BYTES = 10 * 1024 * 1024  # 10MB
#: 추출 결과 상한. 넘으면 잘라서 `truncated=True` 로 알린다.
MAX_TEXT_CHARS = 200_000

#: 확장자 → 추출기 키. content-type 은 브라우저마다 제각각이라 신뢰하지 않는다.
PDF_EXTENSIONS = frozenset({".pdf"})
DOCX_EXTENSIONS = frozenset({".docx"})
TEXT_EXTENSIONS = frozenset({".txt", ".md", ".markdown", ".csv", ".json", ".yaml", ".yml"})

SUPPORTED_EXTENSIONS = PDF_EXTENSIONS | DOCX_EXTENSIONS | TEXT_EXTENSIONS


class UnsupportedDocumentError(ValueError):
    """지원하지 않는 확장자."""


class DocumentTooLargeError(ValueError):
    """`MAX_UPLOAD_BYTES` 초과."""


class DocumentParseError(ValueError):
    """파일이 손상됐거나 확장자와 실제 내용이 다르다."""


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    text: str
    chars: int
    pages: int | None
    truncated: bool


def _extension(filename: str) -> str:
    idx = filename.rfind(".")
    return filename[idx:].lower() if idx != -1 else ""


def _extract_pdf(data: bytes) -> tuple[str, int]:
    import pdfplumber

    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            # 텍스트 레이어가 없는(스캔) 페이지는 None 을 준다 — OCR 은 범위 밖이다.
            parts.append(page.extract_text() or "")
        pages = len(pdf.pages)
    return "\n\n".join(parts), pages


def _extract_docx(data: bytes) -> tuple[str, None]:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    # 표 안의 글자도 원문의 일부다. 빼면 사용자는 "내용이 사라졌다" 로 본다.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts), None


def _extract_text(data: bytes) -> tuple[str, None]:
    for encoding in ("utf-8", "utf-16", "cp949", "latin-1"):
        try:
            return data.decode(encoding), None
        except (UnicodeDecodeError, LookupError):
            continue
    raise DocumentParseError("could not decode as text")


def extract_document(filename: str, data: bytes) -> ExtractedDocument:
    """파일 바이트 → 텍스트. 실패는 전부 이 모듈의 예외 3종으로 좁혀 던진다."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentTooLargeError(f"{len(data)} bytes > {MAX_UPLOAD_BYTES}")

    ext = _extension(filename)
    if ext in PDF_EXTENSIONS:
        extractor = _extract_pdf
    elif ext in DOCX_EXTENSIONS:
        extractor = _extract_docx
    elif ext in TEXT_EXTENSIONS:
        extractor = _extract_text
    else:
        raise UnsupportedDocumentError(ext or filename)

    try:
        raw, pages = extractor(data)
    except (UnsupportedDocumentError, DocumentTooLargeError, DocumentParseError):
        raise
    except Exception as exc:  # noqa: BLE001 — 라이브러리별 예외가 제각각이라 한 종류로 좁힌다
        logger.debug("document extraction failed filename=%s", filename, exc_info=True)
        raise DocumentParseError(str(exc)) from exc

    # 페이지마다 빈 줄이 끼어 세 줄 이상 이어지는 것만 정리한다. 문단 구분(빈 줄 1개)은
    # 원문의 정보라 보존한다 — 이 텍스트는 그대로 프롬프트에 들어간다.
    text = "\n".join(line.rstrip() for line in raw.splitlines()).strip()
    while "\n\n\n" in text:
        text = text.replace("\n\n\n", "\n\n")

    truncated = len(text) > MAX_TEXT_CHARS
    if truncated:
        text = text[:MAX_TEXT_CHARS]

    return ExtractedDocument(
        filename=filename, text=text, chars=len(text), pages=pages, truncated=truncated,
    )


__all__ = [
    "DOCX_EXTENSIONS",
    "MAX_TEXT_CHARS",
    "MAX_UPLOAD_BYTES",
    "PDF_EXTENSIONS",
    "SUPPORTED_EXTENSIONS",
    "TEXT_EXTENSIONS",
    "DocumentParseError",
    "DocumentTooLargeError",
    "ExtractedDocument",
    "UnsupportedDocumentError",
    "extract_document",
]
